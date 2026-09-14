import io
import os
from typing import Tuple, Dict, Any

# Import text extraction libraries with fallback support
try:
    import fitz  # PyMuPDF
    HAS_FITZ = True
except Exception:
    HAS_FITZ = False

try:
    import pdfplumber
    HAS_PDFPLUMBER = True
except Exception:
    HAS_PDFPLUMBER = False

try:
    import docx
    HAS_DOCX = True
except Exception:
    HAS_DOCX = False

MAX_FILE_SIZE_BYTES = 15 * 1024 * 1024  # 15 MB limit

class ResumeParsingError(Exception):
    pass

def extract_text_from_pdf(file_bytes: bytes) -> str:
    text_content = ""
    # Method 1: Try PyMuPDF (fitz)
    if HAS_FITZ:
        try:
            doc = fitz.open(stream=file_bytes, filetype="pdf")
            for page in doc:
                page_text = page.get_text()
                if page_text:
                    text_content += page_text + "\n"
            if text_content.strip():
                return text_content.strip()
        except Exception as e:
            print(f"[ResumeParser] fitz extraction failed: {e}")

    # Method 2: Fallback to pdfplumber
    if HAS_PDFPLUMBER:
        try:
            with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
                for page in pdf.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text_content += page_text + "\n"
            if text_content.strip():
                return text_content.strip()
        except Exception as e:
            print(f"[ResumeParser] pdfplumber extraction failed: {e}")

    if not text_content.strip():
        raise ResumeParsingError("Unable to extract text from the PDF file. The PDF may be empty, image-scanned, or corrupted.")

    return text_content.strip()

def extract_text_from_docx(file_bytes: bytes) -> str:
    if not HAS_DOCX:
        raise ResumeParsingError("DOCX processing library (python-docx) is not installed.")
    try:
        doc = docx.Document(io.BytesIO(file_bytes))
        full_text = []
        for para in doc.paragraphs:
            if para.text.strip():
                full_text.append(para.text.strip())
        for table in doc.tables:
            for row in table.rows:
                row_data = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_data:
                    full_text.append(" | ".join(row_data))
        text = "\n".join(full_text)
        if not text.strip():
            raise ResumeParsingError("Extracted DOCX text is empty.")
        return text.strip()
    except Exception as e:
        raise ResumeParsingError(f"Error parsing DOCX file: {str(e)}")

def extract_text_from_txt(file_bytes: bytes) -> str:
    try:
        # Try UTF-8 first, fallback to latin-1
        try:
            text = file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            text = file_bytes.decode("latin-1")
        if not text.strip():
            raise ResumeParsingError("TXT file is empty.")
        return text.strip()
    except Exception as e:
        raise ResumeParsingError(f"Error decoding TXT file: {str(e)}")

def parse_resume_bytes(filename: str, file_bytes: bytes) -> Tuple[str, str]:
    """
    Main parser router. Returns (extracted_text, file_type).
    Validates file size, extension, and non-empty content.
    """
    if not file_bytes:
        raise ResumeParsingError("File is empty (0 bytes).")

    if len(file_bytes) > MAX_FILE_SIZE_BYTES:
        raise ResumeParsingError(f"File size exceeds maximum allowed limit of {MAX_FILE_SIZE_BYTES / (1024*1024)}MB.")

    ext = os.path.splitext(filename)[1].lower()

    if ext == ".pdf":
        text = extract_text_from_pdf(file_bytes)
        return text, "PDF"
    elif ext in [".docx", ".doc"]:
        text = extract_text_from_docx(file_bytes)
        return text, "DOCX"
    elif ext in [".txt", ".md"]:
        text = extract_text_from_txt(file_bytes)
        return text, "TXT"
    else:
        raise ResumeParsingError(f"Unsupported file format '{ext}'. Supported formats: PDF, DOCX, TXT.")
